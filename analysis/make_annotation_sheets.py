"""Build the tagger hand-check sheets (one Excel file per annotator).

100 LLM sentence units, drawn evenly across the 12 conditions (seeded): 20 SHARED units that
every annotator labels (to measure human-human agreement) + 20 OWN units per annotator (4
annotators). Each annotator gets 40 rows in random order, so shared and own items look the
same. Each row shows the two previous units for context; the tagger's answer is hidden.

The answer key (tagger labels, which items are shared) goes to annotation/_key.csv, which is
gitignored, so annotators can't see it. Score the filled sheets with
analysis/score_annotation.py.

    python analysis/make_annotation_sheets.py --cache <dialogtag cache json>
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import pathlib
import random
import sys

ROOT = pathlib.Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from analysis.dialogue_acts import COARSE_LABELS, FINE_TO_COARSE, load_generated  # noqa: E402

N_ANNOTATORS, N_SHARED, N_OWN, SEED = 4, 20, 20, 20261005
OUT = ROOT / "annotation"

GUIDE = [
    ("Statement", "Describes a fact, an event or the speaker's own experience. No evaluation.",
     "I live in Dallas. / We drove to the lake last summer. / My son plays soccer."),
    ("Opinion", "The speaker's view, judgement or belief: 'I think…', 'it's important…', 'X is better'.",
     "I think nursing homes can be a good option. / Public schools are underfunded."),
    ("Backchannel", "A short listener reaction with no new content: acknowledging or appreciating.",
     "Uh-huh. / Yeah. / Oh, wow. / I see. / That's great! / That sounds nice."),
    ("YesNoQuestion", "A question answerable with yes/no, including tag questions.",
     "Do you have pets? / Have you ever been camping? / It's hot there, isn't it?"),
    ("WhQuestion/OpenQuestion", "An open question (what/how/why/where…), including 'How about you?'.",
     "What kind of music do you like? / How about you? / What do you think?"),
    ("Agreement", "Accepting or agreeing with what the other person just said.",
     "I agree. / Exactly. / That's true. / Absolutely. / You're right."),
    ("Answer", "A direct, short answer to a question: yes / no / don't know.",
     "Yes. / No, I haven't. / I'm not sure."),
    ("Directive", "Asks or suggests the other person do something, or offers to do something.",
     "You should try it. / Let me know what you think. / Let's talk about cars."),
    ("Repair/Hedge", "Hedging, rephrasing what was said, checking understanding, holding the turn.",
     "Well, I mean… / So you're saying that… / Sorry, what was that?"),
    ("Abandoned/Other", "Cut-off sentences, greetings, thanks, goodbyes, apologies, anything else.",
     "Hi! / Thanks for chatting! / Have a great day. / Bye. / So I was going to—"),
]


def _key(text: str) -> str:
    return hashlib.sha1(text.encode("utf-8")).hexdigest()


def sample_units(cache_path: pathlib.Path) -> list[dict]:
    memo = json.load(open(cache_path, encoding="utf-8"))["generated_unit_labels"]
    rng = random.Random(SEED)
    by_cond: dict[str, list[dict]] = {}
    for conv in load_generated(ROOT / "data" / "generated_v3"):
        for i, text in enumerate(conv.texts):
            if _key(text) not in memo or len(text.split()) < 1:
                continue
            ctx = [f"{conv.speakers[j][-1]}: {conv.texts[j]}" for j in range(max(0, i - 2), i)]
            by_cond.setdefault(conv.condition, []).append({
                "condition": conv.condition, "conversation_no": conv.conversation_no, "unit_index": i,
                "speaker": conv.speakers[i][-1], "context": "\n".join(ctx), "sentence": text,
                "tagger_fine": memo[_key(text)], "tagger_coarse": FINE_TO_COARSE[memo[_key(text)]],
            })
    conds = sorted(by_cond)
    total = N_SHARED + N_ANNOTATORS * N_OWN
    quota = {c: total // len(conds) for c in conds}
    for c in rng.sample(conds, total - sum(quota.values())):     # spread the remainder
        quota[c] += 1
    picked = [u for c in conds for u in rng.sample(by_cond[c], quota[c])]
    rng.shuffle(picked)
    return picked


def write_sheet(path: pathlib.Path, annotator: int, rows: list[dict]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Label these"
    head = ["#", "Context (previous lines)", "Speaker", "SENTENCE TO LABEL", "Your label", "Note (optional)"]
    ws.append(head)
    for r, row in enumerate(rows, start=1):
        ws.append([r, row["context"], row["speaker"], row["sentence"], "", ""])
    for col, width in zip("ABCDEF", (5, 60, 9, 60, 26, 30)):
        ws.column_dimensions[col].width = width
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        row[1].font = Font(color="808080")
        row[3].font = Font(bold=True)
        row[4].fill = PatternFill("solid", fgColor="FFF2CC")
    dv = DataValidation(type="list", formula1='"' + ",".join(COARSE_LABELS) + '"', allow_blank=True)
    dv.add(f"E2:E{len(rows) + 1}")
    ws.add_data_validation(dv)
    ws.freeze_panes = "A2"

    g = wb.create_sheet("Label guide")
    g.append([f"Annotator {annotator}: pick ONE label per sentence from the dropdown (yellow column)."])
    g.append(["Label the SENTENCE in bold; the grey context is only there to help. Don't look at "
              "other people's answers. If unsure, pick the best fit and write a note."])
    g.append([])
    g.append(["Label", "What it means", "Examples"])
    for item in GUIDE:
        g.append(list(item))
    for col, width in zip("ABC", (26, 70, 70)):
        g.column_dimensions[col].width = width
    for cell in g[4]:
        cell.font = Font(bold=True)
    for row in g.iter_rows(min_row=5):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    g["A1"].font = Font(bold=True, size=12)
    wb.save(path)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--cache", required=True, type=pathlib.Path)
    args = ap.parse_args()
    units = sample_units(args.cache)
    for i, u in enumerate(units, start=1):
        u["item_id"] = f"U{i:03d}"
    shared = units[:N_SHARED]
    own = [units[N_SHARED + a * N_OWN: N_SHARED + (a + 1) * N_OWN] for a in range(N_ANNOTATORS)]

    (OUT / "sheets").mkdir(parents=True, exist_ok=True)
    rng = random.Random(SEED + 1)
    key_rows = []
    for a in range(N_ANNOTATORS):
        rows = shared + own[a]
        rng.shuffle(rows)
        write_sheet(OUT / "sheets" / f"annotator_{a + 1}.xlsx", a + 1, rows)
        for r, u in enumerate(rows, start=1):
            key_rows.append({"annotator": a + 1, "row": r, "item_id": u["item_id"],
                             "shared": int(u in shared), "condition": u["condition"],
                             "conversation_no": u["conversation_no"], "unit_index": u["unit_index"],
                             "tagger_coarse": u["tagger_coarse"], "tagger_fine": u["tagger_fine"]})
    with open(OUT / "_key.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(key_rows[0]))
        w.writeheader()
        w.writerows(key_rows)
    conds = sorted({u["condition"] for u in units})
    print(f"{len(units)} units ({N_SHARED} shared + {N_ANNOTATORS}x{N_OWN} own) from {len(conds)} conditions;"
          f" sheets in {OUT / 'sheets'}; key in {OUT / '_key.csv'} (gitignored)")
    print("per condition:", {c: sum(u['condition'] == c for u in units) for c in conds})
    print("tagger labels in sample:", {l: sum(u['tagger_coarse'] == l for u in units) for l in COARSE_LABELS})


if __name__ == "__main__":
    main()
